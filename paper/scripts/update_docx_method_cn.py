#!/usr/bin/env python3
"""Apply the reviewed Chinese Method revision to the Word draft.

The script preserves existing OMML equations instead of recreating them.
"""

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PAPER_DIR = Path(__file__).resolve().parents[1]
SOURCE = PAPER_DIR / "ICRA2026.docx"
BACKUP = PAPER_DIR / "ICRA2026_before_method_revision.docx"


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def append_text(paragraph_element, text):
    if not text:
        return
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    paragraph_element.append(run)


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text):
    clear_content(paragraph)
    append_text(paragraph._p, text)


def set_mixed(paragraph, parts, math_cache):
    """Replace text while inserting cloned equations as (paragraph_index, math_index)."""
    clear_content(paragraph)
    for part in parts:
        if isinstance(part, str):
            append_text(paragraph._p, part)
        else:
            paragraph._p.append(deepcopy(math_cache[part[0]][part[1]]))


def top_level_math(paragraph):
    return [
        child
        for child in paragraph._p
        if child.tag in (qn("m:oMath"), qn("m:oMathPara"))
    ]


def replace_math_symbol(paragraph, old, new):
    for node in paragraph._p.iter(qn("m:t")):
        if node.text == old:
            node.text = new


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    # Rebuild from the untouched backup on repeated runs so paragraph indices
    # and preserved equation objects remain deterministic.
    document = Document(BACKUP)
    paragraphs = list(document.paragraphs)

    if paragraphs[48].text.strip() != "PRELIMINARIES":
        raise RuntimeError("Unexpected document structure at PRELIMINARIES")
    if paragraphs[58].text.strip() != "Method":
        raise RuntimeError("Unexpected document structure at Method")

    # Avoid overloading theta: phi denotes the polar angle, while theta remains
    # reserved for camera model parameters.
    for index in (91, 92, 93):
        replace_math_symbol(paragraphs[index], "θ", "φ")

    math_cache = {
        index: [deepcopy(node) for node in top_level_math(paragraphs[index])]
        for index in range(59, 99)
    }

    # PRELIMINARIES is temporarily removed. The contribution draft below it is
    # retained because it belongs to the unfinished Introduction material.
    remove_paragraph(paragraphs[49])
    remove_paragraph(paragraphs[48])
    # These were empty PRELIMINARIES placeholders and would otherwise become
    # orphaned subsections after the section heading is removed.
    remove_paragraph(paragraphs[57])
    remove_paragraph(paragraphs[56])

    set_plain(paragraphs[59], "Multi-Board Target Formulation")
    set_mixed(
        paragraphs[60],
        [
            "本文使用的 multi-board calibration target 由多个具有唯一 AprilTag ID 的 board 组成，记为 ",
            (60, 0),
            "。其中 ",
            (60, 1),
            " 为 reference board。第 ",
            (60, 3),
            " 个 board 的局部坐标系记为 ",
            (60, 4),
            "，其第 ",
            (60, 5),
            " 个三维控制点记为 ",
            (60, 6),
            "。控制点包括 tag 的 outer corners 和本文生成的 internal points。",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[61],
        [
            "我们以 ",
            (61, 0),
            " 作为多板目标的公共坐标系，并将 reference board 的变换固定为单位变换。记 ",
            (61, 1),
            " 为从 ",
            (61, 2),
            " 到 ",
            (61, 3),
            " 的刚体变换。对于第 ",
            (61, 4),
            " 帧，相机坐标系记为 ",
            (61, 5),
            "，从 reference board 到相机坐标系的变换记为 ",
            (61, 6),
            "。因此，board ",
            (61, 7),
            " 上的点 ",
            (61, 8),
            " 在第 ",
            (61, 9),
            " 帧相机坐标系下为",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[63],
        ["给定相机模型参数 ", (63, 0), " 及投影函数 ", (63, 1), "，其预测图像位置为"],
        math_cache,
    )
    set_mixed(
        paragraphs[65],
        [
            "对应的图像观测记为 ",
            (65, 0),
            "。所有有效观测构成集合 ",
            (65, 1),
            "，每项包含 frame、board 和 point 索引及二维坐标。该表示区分三类待估变量：相机模型参数、各帧的 reference-board pose，以及其余 board 相对于 reference board 的刚体变换。",
        ],
        math_cache,
    )

    set_plain(paragraphs[66], "Camera Model Initialization")
    set_mixed(
        paragraphs[67],
        [
            "标定开始时尚不存在可靠的相机模型，而良好的模型初值对宽视场优化尤为重要。多板 AprilTag target 覆盖较大视场，其 outer corners 可由标签边界稳定定位。我们因此使用这些观测初始化相机模型。对于给定的模型族，初始参数 ",
            (67, 0),
            " 由图像分辨率和模型参数约束构造，不读取外部标定文件或预设内参。",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[68],
        [
            "给定第 ",
            (68, 0),
            " 帧中 board ",
            (68, 1),
            " 的 outer-corner 集合 ",
            (68, 2),
            "，其三维坐标和图像观测分别为 ",
            (68, 3),
            " 与 ",
            (68, 4),
            "。在 ",
            (68, 5),
            " 下，该 frame-board pose 由下式初始化：",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[70],
        [
            "我们在多帧和多 board 观测上执行该初始化，并剔除位姿无效或重投影误差过大的观测。随后以 ",
            (70, 0),
            " 和 ",
            (70, 1),
            " 为初值，优化相机参数及所有有效 frame-board poses：",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[72],
        [
            "其中 ",
            (72, 0),
            " 为鲁棒损失。该阶段不引入未知的 board 间结构；每个有效 frame-board observation 具有独立位姿。优化得到的 ",
            (72, 1),
            " 初始化后续完整标定，并在下一节作为 intermediate camera model。",
        ],
        math_cache,
    )

    set_plain(paragraphs[73], "Internal Point Generation")
    set_plain(
        paragraphs[74],
        "Outer corners 能够稳定定位 board 边界，但每个 tag 仅提供四个观测，难以充分约束大视场区域。我们利用初始化后的相机模型建立图像点与 viewing ray 的映射，并为 board lattice 上的 internal points 生成可精化的初始位置。",
    )
    set_mixed(
        paragraphs[75],
        [
            "给定 intermediate camera parameters ",
            (75, 0),
            "，对于第 ",
            (75, 1),
            " 帧中的 board ",
            (75, 2),
            "，我们沿用 outer-corner pose initialization 得到 ",
            (75, 3),
            "。对于 internal-point 集合 ",
            (75, 4),
            " 中的点 ",
            (75, 5),
            "，该模型首先给出名义几何预测，作为边界模型不可用时的回退。",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[77],
        [
            "主要 seed 由 board 边界在 viewing-ray domain 中确定。我们沿四条边缘采样支持点，经 intermediate camera model 反投影至 ",
            (77, 0),
            "，并拟合上、下、左、右四条球面边界 ",
            (77, 1),
            "、",
            (77, 2),
            "、",
            (77, 3),
            " 和 ",
            (77, 4),
            "。对内部点 ",
            (77, 5),
            "，令 ",
            (77, 6),
            " 表示其在 board lattice 水平和竖直方向上的归一化坐标。上下与左右边界分别通过 spherical interpolation 得到",
        ],
        math_cache,
    )
    set_plain(paragraphs[79], "两条插值射线归一化后形成 seed ray，并投影回原图：")
    set_mixed(
        paragraphs[81],
        [
            "我们以 ",
            (81, 0),
            " 锚定局部球面搜索，在 viewing-ray domain 中调整候选位置，再于原始图像的局部邻域执行 subpixel refinement。",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[83],
        [
            "若 refined point 超出图像边界、局部图像证据不足或相对 seed 的位移超过阈值，则将其标记为无效。其余 internal points 保留 board 和 lattice point 索引，并加入观测集合 ",
            (65, 1),
            "，与 outer corners 一同进入后端优化。",
        ],
        math_cache,
    )
    remove_paragraph(paragraphs[76])
    remove_paragraph(paragraphs[82])

    set_plain(paragraphs[84], "Spherical Bundle Adjustment")
    set_plain(
        paragraphs[85],
        "标准重投影误差在图像平面上度量投影点与观测点的距离。对于宽视场相机，相同像素误差在不同视场区域并不对应相同的射线偏差，尤其在高 polar angle 区域。为此，我们在 viewing-ray domain 中引入球面约束。",
    )
    set_mixed(
        paragraphs[86],
        [
            "给定观测点 ",
            (86, 0),
            "，我们使用当前相机参数将其反投影至单位球面，并归一化为观测射线 ",
            (86, 1),
            "；该射线随优化中的相机参数更新。对应控制点 ",
            (86, 2),
            " 经上一节定义的多板变换链进入相机坐标系，并归一化为预测射线：",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[88],
        [
            "在观测射线处构造局部切平面基 ",
            (88, 0),
            "，其两列与 ",
            (88, 1),
            " 正交。tangent-plane residual 定义为",
        ],
        math_cache,
    )
    set_plain(
        paragraphs[90],
        "该二维残差保留局部射线偏差的两个自由度，可直接用于最小二乘优化。",
    )
    set_mixed(
        paragraphs[91],
        [
            "为在中心区域保留稳定的像素约束，并逐步增强高 polar angle 区域的射线约束，我们根据 polar angle ",
            (91, 0),
            " 定义连续权重：",
        ],
        math_cache,
    )
    set_mixed(
        paragraphs[93],
        [
            "其中 ",
            (93, 0),
            " 为 sigmoid 函数，",
            (93, 1),
            " 为过渡角度，",
            (93, 2),
            " 控制过渡宽度。权重随 polar angle 增大而提高。",
        ],
        math_cache,
    )
    set_plain(
        paragraphs[94],
        "像素残差与 tangent-plane residual 的量纲不同。我们使用初始化焦距几何均值定义固定参考尺度",
    )
    set_plain(paragraphs[96], "并将球面残差缩放至近似像素量级。最终的 polar-aware hybrid objective 为")
    set_mixed(
        paragraphs[98],
        [
            "其中 ",
            (98, 0),
            " 和 ",
            (98, 1),
            " 为对应的鲁棒核。优化变量包括相机参数、各帧 reference-board poses 和 inter-board transforms，并固定 reference board 的 gauge。该目标在中心视场保留像素约束，在高 polar angle 区域增强射线域约束。",
        ],
        math_cache,
    )

    document.save(SOURCE)


if __name__ == "__main__":
    main()
